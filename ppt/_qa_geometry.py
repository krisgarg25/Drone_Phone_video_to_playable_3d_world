"""Geometric QA: detect text hidden beneath an opaque shape drawn later.

That is exactly the bug found by rendering slide 3 (evidence strip drawn
on top of the last table row), so it is checked on every slide.
"""
import pptx
from pptx.util import Emu

OUT = r'c:\Users\krisg\Desktop\Drone to 3d mesh\ppt\SIH26158_ONEPASS_Idea_Deck.pptx'
prs = pptx.Presentation(OUT)


def box(sh):
    return (Emu(sh.left).inches, Emu(sh.top).inches,
            Emu(sh.left).inches + Emu(sh.width).inches,
            Emu(sh.top).inches + Emu(sh.height).inches)


def overlap(a, b):
    x = min(a[2], b[2]) - max(a[0], b[0])
    y = min(a[3], b[3]) - max(a[1], b[1])
    return max(0.0, x) * max(0.0, y)


def is_opaque_fill(sh):
    try:
        if sh.fill.type is None:
            return False
        # solid fill (type 1) on an autoshape/textbox occludes what is below
        return 'SOLID' in str(sh.fill.type)
    except Exception:
        return False


issues = []
for i, s in enumerate(prs.slides):
    n = i + 1
    shapes = list(s.shapes)
    for zi, sh in enumerate(shapes):
        if not sh.has_text_frame:
            continue
        if not sh.text_frame.text.strip():
            continue
        tbx = box(sh)
        # only opaque shapes LATER in z-order can hide this text
        for other in shapes[zi + 1:]:
            if not is_opaque_fill(other):
                continue
            ob = box(other)
            a = overlap(tbx, ob)
            area = (tbx[2] - tbx[0]) * (tbx[3] - tbx[1])
            if area > 0 and a / area > 0.5:
                snip = sh.text_frame.text.replace('\n', ' ')[:44]
                issues.append(
                    'slide %d: %r (%d%% covered by %s)' %
                    (n, snip, round(100 * a / area), other.name))

print('Checked %d slides' % len(prs.slides))
if issues:
    print('HIDDEN TEXT FOUND:')
    for x in issues:
        print('  ' + x)
else:
    print('OK: no text obscured by an opaque shape above it')

# explicit slide-3 geometry proof
print('\nslide 3 geometry:')
s3 = prs.slides[2]
last_row = None
strip = None
for sh in s3.shapes:
    if sh.has_text_frame and 'Interface / QA' in sh.text_frame.text:
        last_row = box(sh)
    if sh.shape_type is not None and 'AUTO_SHAPE' in str(sh.shape_type):
        b = box(sh)
        if abs(b[1] - 5.34) < 0.01 and (b[2] - b[0]) > 12:
            strip = b
print('  Interface / QA row box : %s' %
      (['%.2f' % v for v in last_row] if last_row else 'NOT FOUND'))
print('  evidence strip top     : %.2f' % strip[1] if strip else '  strip NOT FOUND')
if last_row and strip:
    print('  gap (strip_top - row_bottom) = %.2f in  -> %s' %
          (strip[1] - last_row[3],
           'CLEAR' if strip[1] >= last_row[3] else 'OVERLAP'))
