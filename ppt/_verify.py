import pptx
from pptx.util import Emu, Pt

OUT = r'c:\Users\krisg\Desktop\Drone to 3d mesh\ppt\SIH26158_ONEPASS_Idea_Deck.pptx'
prs = pptx.Presentation(OUT)

W = Emu(prs.slide_width).inches
H = Emu(prs.slide_height).inches
print('SIZE %.2f x %.2f in' % (W, H))
print('SLIDE COUNT: %d  (must be 6)\n' % len(prs.slides))

REQUIRED = {
    1: [],
    2: ['Proposed Solution (Describe your Idea/Solution/Prototype)',
        'Detailed explanation of the proposed solution',
        ' How it addresses the problem',
        'Innovation and uniqueness of the solution '],
    3: ['Technologies to be used (e.g. programming languages, frameworks, hardware)',
        'Methodology and process for implementation (Flow Charts/Images/ working prototype)'],
    4: ['Analysis of the feasibility of the idea',
        'Potential challenges and risks',
        'Strategies for overcoming these challenges'],
    5: ['Potential impact on the target audience',
        'Benefits of the solution (social, economic, environmental, etc.)'],
    6: ['Details / Links of the reference and research work'],
}

problems = []
alltext = {}

for i, s in enumerate(prs.slides):
    n = i + 1
    texts = []
    overflow = []
    for sh in s.shapes:
        # bounds check
        try:
            l = Emu(sh.left).inches; t = Emu(sh.top).inches
            r = l + Emu(sh.width).inches; b = t + Emu(sh.height).inches
            if l < -0.05 or t < -0.60 or r > W + 0.05 or b > H + 0.05:
                overflow.append('%s (L%.2f T%.2f R%.2f B%.2f)' % (sh.name, l, t, r, b))
        except Exception:
            pass
        if sh.has_text_frame:
            tx = sh.text_frame.text
            if tx.strip():
                texts.append(tx)
    alltext[n] = '\n'.join(texts)

    print('SLIDE %d: %2d shapes, %2d with text' % (n, len(s.shapes),
                                                    len(texts)))
    if overflow:
        print('   OUT OF BOUNDS: %s' % overflow)
        problems.append('slide %d bounds' % n)

    # pointer check
    for req in REQUIRED.get(n, []):
        if req not in alltext[n]:
            print('   MISSING POINTER: %r' % req)
            problems.append('slide %d pointer %r' % (n, req[:35]))

    # text-wall check: any single shape with >6 lines
    for sh in s.shapes:
        if sh.has_text_frame:
            lines = sh.text_frame.text.count('\n') + 1
            if lines > 6:
                print('   TEXT WALL (%d lines): %s' % (lines, sh.name))
                problems.append('slide %d wall %s' % (n, sh.name))

print('\n--- pointer audit ---')
for n, reqs in REQUIRED.items():
    if not reqs:
        continue
    ok = sum(1 for r in reqs if r in alltext[n])
    print('slide %d: %d/%d pointers verbatim' % (n, ok, len(reqs)))

print('\n--- slide 1 field check ---')
for k in ('SIH26158', 'Software', 'Theme', 'Team ID'):
    print('  %-12s %s' % (k, 'present' if k in alltext[1] else 'ABSENT'))

print('\nRESULT:', 'ALL CHECKS PASSED' if not problems else 'ISSUES: %s' % problems)
