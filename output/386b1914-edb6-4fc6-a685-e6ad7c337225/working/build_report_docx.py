"""Stage 3 — build the report DOCX, then export the PDF through Word.

Text lives in the content() function below; layout helpers live above it. Rerun
build_report_docx.py and then export_pdf.py to regenerate both deliverables.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = r"C:\Users\krisg\Desktop\Drone to 3d mesh\output\386b1914-edb6-4fc6-a685-e6ad7c337225"
FIGURE = BASE + r"\stage2\images\methodology_diagram.png"
DOCX = BASE + r"\stage3\Drone-to-3D-Walkable-World-Project-Report.docx"

SERIF = "Times New Roman"
SANS = "Arial"
MONO = "Consolas"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x01, 0x57, 0x9B)
PAGE_W = Cm(16.0)  # A4 (21 cm) minus 2 x 2.5 cm margins

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


# --------------------------------------------------------------- low-level helpers
def set_run_font(run, name=SERIF, size=12, bold=False, italic=False, color=None, mono=False,
                 spacing=None):
    run.font.name = MONO if mono else name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    face = MONO if mono else name
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), face)
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        rPr.append(sp)


def add_runs(par, text, size=12, name=SERIF, color=None, bold_color=None, spacing=None):
    """Render inline **bold**, *italic* and `code` spans."""
    for part in TOKEN.split(text):
        if not part:
            continue
        bold = italic = mono = False
        if part.startswith("**") and part.endswith("**"):
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*"):
            part, italic = part[1:-1], True
        run = par.add_run(part)
        set_run_font(run, name=name, size=size, bold=bold, italic=italic, mono=mono,
                     color=(bold_color if bold and bold_color else color), spacing=spacing)


def para(doc, text="", size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8,
         line=1.5, name=SERIF, color=None, bold_color=None, keep_next=False, spacing=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if text:
        add_runs(p, text, size=size, name=name, color=color, bold_color=bold_color,
                 spacing=spacing)
    return p


def style_font(style, name, size, bold=False, color=INK, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def border(par, edge="bottom", sz=8, color="1F2937", space=4):
    pPr = par._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def cell_border(cell, edge, sz=8, color="1F2937"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    borders.append(el)


def bare_table(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        b.append(el)
    tblPr.append(b)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 60), ("left", 110), ("bottom", 60), ("right", 110)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def three_line_table(doc, rows, widths, size=11):
    """Booktabs-style table: heavy top/bottom rules, light rule under the header."""
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.alignment = 1
    table.autofit = False
    bare_table(table)
    for r, row in enumerate(rows):
        tr = table.rows[r]
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if r == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for c, cell_text in enumerate(row):
            cell = tr.cells[c]
            cell.width = widths[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
            pf.line_spacing = 1.15
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, cell_text, size=size, name=SERIF, bold_color=INK)
            if r == 0:
                for run in p.runs:
                    run.bold = True
                cell_border(cell, "top", sz=16)
                cell_border(cell, "bottom", sz=8)
            elif r == len(rows) - 1:
                cell_border(cell, "bottom", sz=16)
    # keep every row, and the caption that follows, on the same page
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    return table


def add_field(par, instr, placeholder="", size=12, name=SERIF):
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    for r in (par.add_run(), par.add_run(), par.add_run()):
        set_run_font(r, name=name, size=size)
    par.runs[-3]._r.append(fld("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    par.runs[-2]._r.append(it)
    par.runs[-1]._r.append(fld("separate"))
    res = par.add_run(placeholder)
    set_run_font(res, name=name, size=size, color=MUTED, italic=True)
    end = par.add_run()
    set_run_font(end, name=name, size=size)
    end._r.append(fld("end"))


def caption(doc, text):
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=14, line=1.15)
    add_runs(p, text, size=10, name=SANS, color=MUTED)
    return p


def references(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.9)
        pf.space_after = Pt(7)
        pf.line_spacing = 1.25
        add_runs(p, f"[{i}]  " + item, size=11)


# --------------------------------------------------------------------- page chrome
def chrome(doc):
    st = doc.styles
    normal = st["Normal"]
    style_font(normal, SERIF, 12)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.5

    for name, size, before, after in (("Heading 1", 14, 18, 9), ("Heading 2", 12, 11, 5)):
        h = st[name]
        style_font(h, SANS, size, bold=True, color=INK)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.line_spacing = 1.25
        h.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.1)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(2)
    add_runs(hp, "WALKABLE — Turning a Single Drone Video into a Metric, Walkable 3D World",
             size=8.5, name=SANS, color=MUTED)
    border(hp, "bottom", sz=6, color="BFBFBF", space=3)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, " PAGE ", "1", size=9.5, name=SANS)
    add_runs(fp, " of ", size=9.5, name=SANS, color=MUTED)
    add_field(fp, " SECTIONPAGES ", "1", size=9.5, name=SANS)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


# ---------------------------------------------------------------------------- content
def content(doc):
    t = para(doc, "WALKABLE: Turning a Single Drone or Handheld Video into a Metric, Walkable 3D World",
             size=16, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, line=1.15, name=SANS)
    for run in t.runs:
        run.bold = True
        run.font.color.rgb = INK
    para(doc, "*A reproducible Structure-from-Motion → 3D Gaussian Splatting → browser-physics pipeline "
              "that runs end-to-end on a 6 GB laptop GPU*",
         size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=16, line=1.2, color=MUTED)

    h1(doc, "1. Project Objectives")
    three_line_table(doc, [
        ["#", "Objective"],
        ["O1", "**Single video → walkable 3D world, one command.** Convert an ordinary drone or handheld "
               "video into a photorealistic, navigable 3D world through one fully automated command — no "
               "manual modelling, no cloud dependency."],
        ["O2", "**Metric, measurable reconstruction.** Recover true metric scale and orientation from minimal "
               "operator input (a single speed or height anchor), so distances, areas, and ground surfaces in "
               "the model are physically meaningful."],
        ["O3", "**Walkability, not just visuals.** Derive physics-valid ground and collision geometry from the "
               "neural reconstruction, enabling first-person walking and autonomous agent navigation in the "
               "scanned scene."],
        ["O4", "**Falsifiable quality on consumer hardware.** Run every stage on a 6 GB laptop GPU and gate "
               "each build on automated acceptance tests (blind visual review, scripted walk tests), not "
               "subjective judgment."],
    ], widths=[Cm(1.4), Cm(14.6)])
    caption(doc, "Table 1. Project objectives for the selected idea: single video → metric, walkable 3D world.")

    h1(doc, "2. Impact Statement")
    para(doc,
         "This work lowers the cost of producing a *usable* 3D replica of a real place from hours of expert "
         "photogrammetry and expensive licensed software to a single command on an ordinary laptop. Because "
         "the output is metric and walkable rather than a view-only render, it directly serves disaster "
         "damage assessment, border and strategic area mapping, construction progress monitoring, "
         "archaeological documentation, urban planning, and mission rehearsal — domains where decisions depend "
         "on measuring and moving through a scene, not merely looking at it. By combining the geometric "
         "guarantees of classical methods with the visual fidelity of neural rendering, and by refusing to "
         "ship any result that fails automated measurement, the pipeline provides a reproducible, low-cost "
         "foundation that researchers, engineers, and field teams can build upon.")

    h1(doc, "3. Research Literature Survey")
    h2(doc, "3.1 Core Reconstruction Methods")
    for t in [
        "**[1] Mildenhall et al. (2020) — NeRF: Representing Scenes as Neural Radiance Fields for View "
        "Synthesis (ECCV 2020).** Stores a scene as a network mapping 5-D coordinates to colour and density, "
        "trained per scene by differentiable volume rendering. Photorealistic, but hours to train and with no "
        "explicit surface — nothing to measure, collide with, or stand on.",
        "**[2] Schönberger & Frahm (2016) — Structure-from-Motion Revisited (COLMAP, CVPR 2016).** The "
        "reference incremental SfM pipeline, and still the accuracy baseline for neural pose estimators. We "
        "adopt it as our pose solver: it fits a 6 GB budget and yields survey-grade registration, at the cost "
        "of runtime and fragility on blur, low texture, and rotation-only motion.",
        "**[3] Kerbl et al. (2023) — 3D Gaussian Splatting for Real-Time Radiance Field Rendering (ACM TOG / "
        "SIGGRAPH 2023).** Represents a scene as 10⁵–10⁶ anisotropic 3-D Gaussians rasterised at interactive "
        "rates, reaching NeRF-class fidelity in minutes on consumer hardware — our appearance stage. Its "
        "known gap, that a splat cloud is not a surface, defines our geometry problem.",
    ]:
        para(doc, t, size=11.5, after=9)

    h2(doc, "3.2 Application-Domain Literature")
    for t in [
        "**[4] Nex & Remondino (2014) — UAV for 3D Mapping Applications: A Review (Applied Geomatics "
        "6(1)).** The classical UAV-photogrammetry baseline — platforms, sensors, flight planning, "
        "georeferencing, and accuracy with and without ground control — and thus the reference against which "
        "our anchor-based metric recovery is positioned.",
        "**[5] Rakha & Gorodetsky (2018) — Review of UAS Applications in the Built Environment (Automation in "
        "Construction 93).** Identifies automated, repeatable capture as the enabler of construction progress "
        "monitoring, which is the basis of our time-series scenario: re-fly the same route, rebuild, and diff "
        "the geometry.",
    ]:
        para(doc, t, size=11.5, after=9)
    para(doc,
         "**Synthesis.** The literature splits into classical photogrammetry (accurate, measurable, slow), "
         "deep SLAM (fast, scale-ambiguous), and neural rendering (photorealistic, surface-free); none of the "
         "three yields a *metric, walkable* world from one video on consumer hardware. Converting "
         "neural-rendering quality into physically usable geometry with verified scale — the gap across "
         "[1]–[5] — is the contribution of this project.", after=14)

    h1(doc, "4. Application Domains")
    for t in [
        "**(i) Border and strategic area mapping.** A single drone pass over remote or sensitive terrain "
        "becomes a measurable 3-D model within minutes on a field laptop, with no cloud link required — "
        "critical where connectivity is denied. Metric terrain supports line-of-sight analysis, route "
        "planning, and change detection between repeat sorties over the same corridor.",
        "**(ii) Disaster damage assessment.** After earthquakes, floods, or landslides, responders can film a "
        "structure or slope and obtain a walkable, measurable model before physical access is safe. Comparing "
        "pre- and post-event models quantifies deformation, debris volume, and accessible paths, while the "
        "walk-test layer verifies that simulated rescue routes are genuinely traversable.",
        "**(iii) Urban planning and smart cities.** Planners gain photorealistic, navigable site models for "
        "design review, shadow and sight-line studies, and public consultation — stakeholders can walk "
        "through a proposal at true scale, and per-site models form the high-fidelity local layers of a "
        "city-scale twin.",
        "**(iv) Construction progress monitoring.** Weekly re-flights of the same route yield comparable 3-D "
        "snapshots: earthwork volumes, structural completeness, and as-built vs. as-planned deviation become "
        "measurable quantities [5]. Because the pipeline is one command and laptop-class, it fits the cadence "
        "and budget of real construction sites.",
        "**(v) Archaeological documentation.** Excavations and heritage structures can be recorded "
        "non-invasively at each dig phase, producing a permanent, measurable 3-D archive even as the site "
        "itself is altered or backfilled [4]. Walkable models support remote scholarship, virtual museum "
        "exhibits, and condition monitoring of fragile monuments.",
        "**(vi) Military reconnaissance and mission planning.** A reconnaissance sortie's footage becomes a "
        "walkable rehearsal environment: teams can virtually traverse the objective area, measure sight lines "
        "and cover from the scanned geometry, and plan insertion routes against a navmesh baked from the real "
        "terrain. All processing is offline and laptop-class, matching denied-communications field "
        "constraints, on Apache/BSD/MIT components with no usage restrictions for this scope.",
    ]:
        para(doc, t, size=11.5, after=9)

    h1(doc, "5. Methodology")
    para(doc, "The methodology runs in three phases, shown in Figure 1.", after=8, keep_next=True)
    fig_p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=0, keep_next=True)
    fig_p.add_run().add_picture(FIGURE, width=PAGE_W)
    caption(doc, "Figure 1. End-to-end methodology of the WALKABLE pipeline: Phase 1 capture and "
                 "pre-processing; Phase 2 3D reconstruction; Phase 3 world building and validation — "
                 "feeding the downstream application domains.")
    for t in [
        "**Phase 1 — Capture and pre-processing.** An ordinary video (drone or handheld), with an optional "
        "ARCore/ARKit pose log, is decoded with OpenCV, scored by variance-of-Laplacian sharpness, and reduced "
        "to a diverse keyframe set. A WebXR capture console colours each surface patch by how many distinct "
        "viewing directions have covered it, so quality is decided at capture time rather than discovered at "
        "training time.",
        "**Phase 2 — 3D reconstruction.** Keyframes are registered by COLMAP, with phone AR tracks injected as "
        "native pose priors that rescue rotation-heavy footage and carry a metric translation into "
        "registration. Those poses seed gsplat training (spherical harmonics degree 3, SSIM+L1, hard Gaussian "
        "cap) within 6 GB of VRAM; orientation comes from gimbal attitude, scale from exactly one named "
        "anchor, and sky and cloud-sea splats are culled by a height-histogram bimodality test.",
        "**Phase 3 — World building and validation.** Splat means are rasterised into a ground heightfield and "
        "a voxelised collision shell exported as GLB; an eleven-assertion quality gate fails the build rather "
        "than shipping a broken world. The validated assets load into a browser runtime (PlayCanvas + "
        "ammo.js/Bullet) for first-person walking and navmesh agents, and acceptance is scored two ways: "
        "blinded A/B visual review against real frames, and scripted traversal in metres walked.",
        "**Design principles.** Each stage is a separate CLI module with typed artefacts on disk, so runs are "
        "resumable and inspectable; re-running a stage invalidates everything downstream; and the neural stage "
        "is trusted with appearance only — never geometry, orientation, or scale.",
    ]:
        para(doc, t, size=11.5, after=9)

    h1(doc, "6. Work Done")
    para(doc,
         "The pipeline is implemented and validated end-to-end: ~18,000 lines of first-party Python 3.12 and "
         "JavaScript across 35 resumable stage modules, driven by one command (`mvp.bat run <scene>`), with a "
         "~550-assertion test suite and a vendored browser runtime that needs no build step. Table 2 lists "
         "measured outcomes on the target hardware (RTX 3050 6 GB laptop).", after=10, keep_next=True)
    three_line_table(doc, [
        ["Result", "Evidence"],
        ["Hardware budget", "Peak training VRAM 1.26 GB at the `high` preset — 5× headroom under the 6 GB card"],
        ["Visual quality", "10/10 blinded A/B pairs judged to be the same scene, no disqualifying artefact"],
        ["Walkability", "65.3 m walked, 16/16 waypoints reached, 0 falls (from 19.9 m, 1/20 before fixes)"],
        ["Ground correctness", "Viewer vs. physics ground discrepancy reduced from 0.101 m to 0.000 m"],
        ["Metric priors", "Synthetic 12 m AR walk reconstructs at scale 1.0000, 7 mm max camera deviation"],
    ], widths=[Cm(4.6), Cm(11.4)], size=10.5)
    caption(doc, "Table 2. Measured outcomes of the implemented pipeline (target hardware: RTX 3050 6 GB laptop).")
    para(doc,
         "Eleven engineering iterations are documented with the measurement that changed each decision, and 2 "
         "of 8 test scenes still fail the route gate — the build reports this rather than hiding it.",
         after=14)

    h1(doc, "References")
    references(doc, [
        "Mildenhall, B., Srinivasan, P. P., Tancik, M., Barron, J. T., Ramamoorthi, R., & Ng, R. (2020). "
        "NeRF: Representing scenes as neural radiance fields for view synthesis. In *European Conference on "
        "Computer Vision (ECCV)*, 405–421.",
        "Schönberger, J. L., & Frahm, J.-M. (2016). Structure-from-motion revisited. In *IEEE Conference on "
        "Computer Vision and Pattern Recognition (CVPR)*, 4104–4113.",
        "Kerbl, B., Kopanas, G., Leimkühler, T., & Drettakis, G. (2023). 3D Gaussian splatting for real-time "
        "radiance field rendering. *ACM Transactions on Graphics*, 42(4), Article 139.",
        "Nex, F., & Remondino, F. (2014). UAV for 3D mapping applications: A review. *Applied Geomatics*, "
        "6(1), 1–15.",
        "Rakha, T., & Gorodetsky, A. (2018). Review of unmanned aerial system (UAS) applications in the built "
        "environment: Towards automated building inspection procedures using drones. *Automation in "
        "Construction*, 93, 252–264.",
    ])


def build():
    doc = Document()
    chrome(doc)
    content(doc)
    doc.save(DOCX)
    print("saved", DOCX)


if __name__ == "__main__":
    if sys.version_info[:2] != (3, 12):
        raise SystemExit("Run with Python 3.12 (has python-docx + matplotlib).")
    build()
